# -*- coding: utf-8 -*-
"""网络管理（IPAM）蓝图：子网管理 / IP 地址管理 / VLAN 管理。"""
import ipaddress as ipmod
import io

from flask import (
    Blueprint, request, jsonify, render_template, redirect, url_for, flash,
    send_file, current_app, session,
)
from flask_login import login_required, current_user

from extensions import db
from models.network_models import Subnet, IPAddress, Vlan
from models.device_group_models import DeviceGroup
from models.models import Device
from forms import SubnetForm, IPAddressForm, VlanForm
from utils.permission import permission_required
from utils.audit import log_audit

network_bp = Blueprint('network', __name__, url_prefix='/network')

IP_STATUS_LABELS = {
    'available': '空闲',
    'reserved': '预留',
    'allocated': '已分配',
    'disabled': '禁用',
}

MAX_AUTO_GENERATE = 4096
MAX_CLUSTER_PREFIX = 22  # 聚类时合并后子网不小于该前缀（/22 约 1024 地址）


# ==================== 工具函数 ====================

def _normalize_ip(value):
    """规范化 IP 字符串，非法返回 None。"""
    if not value:
        return None
    try:
        return str(ipmod.ip_address(str(value).strip()))
    except ValueError:
        return None


def _parse_subnet(network, prefix):
    """解析子网，返回 IPv4Network 或 None。"""
    try:
        net = ipmod.IPv4Network(f'{network}/{prefix}', strict=False)
        return net
    except (ValueError, TypeError):
        return None


def _subnet_overlap(network_obj, exclude_id=None):
    """检查与已有子网是否重叠。"""
    for sub in Subnet.query.filter(Subnet.status != 'disabled'):
        if exclude_id and sub.id == exclude_id:
            continue
        try:
            other = ipmod.IPv4Network(f'{sub.network}/{sub.prefix_length}', strict=False)
        except ValueError:
            continue
        if network_obj.overlaps(other):
            return sub
    return None


def _generate_subnet_ips(subnet, limit=MAX_AUTO_GENERATE):
    """批量生成子网内可用 IP 记录，跳过已存在的。"""
    net = _parse_subnet(subnet.network, subnet.prefix_length)
    if not net:
        return {'created': 0, 'skipped': 0, 'limited': False, 'total': 0}

    existing = {ip.ip_address for ip in subnet.ip_addresses}
    hosts = list(net.hosts())
    total = len(hosts)
    limited = total > limit
    targets = hosts[:limit] if limited else hosts

    created = 0
    skipped = 0
    for host in targets:
        addr = str(host)
        if addr in existing:
            skipped += 1
            continue
        db.session.add(IPAddress(subnet_id=subnet.id, ip_address=addr, status='available'))
        created += 1
    db.session.commit()
    return {'created': created, 'skipped': skipped, 'limited': limited, 'total': total}


def _sync_devices_for_subnet(subnet):
    """把落在子网内的设备 IP 同步为已分配状态。"""
    net = _parse_subnet(subnet.network, subnet.prefix_length)
    if not net:
        return {'created': 0, 'updated': 0, 'matched': 0}

    created = updated = matched = 0
    devices = Device.query.filter(Device.is_decommissioned.is_(False)).all()
    for dev in devices:
        ip = _normalize_ip(dev.management_ip or dev.ip_address)
        if not ip:
            continue
        try:
            if ipmod.ip_address(ip) not in net:
                continue
        except ValueError:
            continue
        matched += 1
        record = IPAddress.query.filter_by(subnet_id=subnet.id, ip_address=ip).first()
        if record:
            if record.status != 'allocated' or record.device_id != dev.id:
                record.status = 'allocated'
                record.device_id = dev.id
                record.hostname = dev.name
                record.mac_address = dev.mac_address
                updated += 1
        else:
            db.session.add(IPAddress(
                subnet_id=subnet.id,
                ip_address=ip,
                status='allocated',
                device_id=dev.id,
                hostname=dev.name,
                mac_address=dev.mac_address,
                description='来自设备库同步',
            ))
            created += 1
    db.session.commit()
    return {'created': created, 'updated': updated, 'matched': matched}


def _collect_group_devices(group):
    """递归收集分组及其所有子分组下的设备（按 id 去重）。"""
    devices = list(group.devices) if group.devices else []
    seen = {d.id for d in devices}
    for child in group.children or []:
        for d in _collect_group_devices(child):
            if d.id not in seen:
                seen.add(d.id)
                devices.append(d)
    return devices


def _device_ip_pairs(devices):
    """提取设备管理/业务 IP，去重后返回 [(ip, device), ...]。"""
    pairs = []
    seen = set()
    for d in devices:
        for raw in (d.management_ip, d.ip_address):
            ip = _normalize_ip(raw)
            if ip and ip not in seen:
                seen.add(ip)
                pairs.append((ip, d))
    return pairs


def _covering_network(ips):
    """计算覆盖一组 IP 的最小子网。ips 为 IPv4Address 列表。"""
    if not ips:
        return None
    first = min(ips)
    last = max(ips)
    for prefix in range(32, -1, -1):
        mask_int = (~((1 << (32 - prefix)) - 1)) & 0xFFFFFFFF if prefix > 0 else 0
        if (int(first) & mask_int) == (int(last) & mask_int):
            return ipmod.IPv4Network((int(first) & mask_int, prefix), strict=False)
    return None


def _cluster_subnet_candidates(ip_list):
    """把 IP 聚成若干候选子网，返回候选列表（按网络地址排序）。"""
    if not ip_list:
        return []
    addr_pairs = sorted(ip_list, key=lambda p: int(ipmod.ip_address(p[0])))
    clusters = []
    for ip, dev in addr_pairs:
        addr = ipmod.ip_address(ip)
        if not clusters:
            clusters.append([(addr, dev)])
            continue
        merged = _covering_network([a for a, _ in clusters[-1]] + [addr])
        if merged and merged.prefixlen >= MAX_CLUSTER_PREFIX:
            clusters[-1].append((addr, dev))
        else:
            clusters.append([(addr, dev)])

    candidates = []
    for cluster in clusters:
        addrs = [a for a, _ in cluster]
        net = _covering_network(addrs)
        if not net:
            continue
        # 单主机/双主机地址退化为最小可用子网 /30
        prefix = min(net.prefixlen, 30)
        network = ipmod.IPv4Network((int(net.network_address), prefix), strict=False)
        devices = [{'name': (dev.name if dev else str(addr)), 'ip': str(addr)}
                   for addr, dev in sorted(cluster, key=lambda p: int(ipmod.ip_address(p[0])))]
        gateway = str(network.network_address + 1) if network.prefixlen <= 30 else ''
        candidates.append({
            'network': str(network.network_address),
            'prefix': prefix,
            'cidr': f'{network.network_address}/{prefix}',
            'netmask': str(network.netmask),
            'broadcast': str(network.broadcast_address),
            'gateway': gateway,
            'count': len(cluster),
            'devices': devices,
        })
    candidates.sort(key=lambda c: (ipmod.ip_address(c['network']), c['prefix']))
    return candidates


def _build_generated_subnet_preset(candidates, index, source_desc, group_name=''):
    """根据候选子网生成新增表单的预填内容。"""
    pick = candidates[min(index, len(candidates) - 1)]
    device_desc = '、'.join(f"{d['name']}({d['ip']})" for d in pick['devices'][:10])
    if len(pick['devices']) > 10:
        device_desc += f" 等共 {pick['count']} 台设备"
    name = f"{group_name}子网" if group_name else f"{pick['cidr']}子网"
    return {
        'name': name,
        'network': pick['network'],
        'prefix_length': pick['prefix'],
        'gateway': pick['gateway'],
        'description': f"由{source_desc}生成（{pick['count']} 台设备，IP：{device_desc}）",
    }


def _export_ips_to_excel(rows, sheet_title='IP地址'):
    """导出 IP 记录到 Excel。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    headers = ['IP地址', '状态', '子网', 'VLAN', '关联设备', '主机名', 'MAC地址', '使用人', '部门', '用途/备注']
    ws.append(headers)
    for ip in rows:
        ws.append([
            ip.ip_address,
            IP_STATUS_LABELS.get(ip.status, ip.status),
            ip.subnet.cidr if ip.subnet else '',
            f"{ip.subnet.vlan.vlan_id} - {ip.subnet.vlan.name}" if ip.subnet and ip.subnet.vlan else '',
            ip.device.name if ip.device else '',
            ip.hostname or '',
            ip.mac_address or '',
            ip.owner or '',
            ip.department or '',
            ip.description or '',
        ])
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col).font = Font(bold=True)
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _send_excel(output, filename):
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def _audit(action, resource_type, resource_id, message, **kwargs):
    log_audit(
        action, resource_type, resource_id, message,
        user_id=current_user.id if current_user.is_authenticated else None,
        **kwargs,
    )


# ==================== 网络管理概览 ====================

@network_bp.route('/')
@login_required
@permission_required('network:view')
def overview():
    subnet_count = Subnet.query.count()
    vlan_count = Vlan.query.count()
    ip_total = IPAddress.query.count()
    ip_allocated = IPAddress.query.filter_by(status='allocated').count()
    ip_available = IPAddress.query.filter_by(status='available').count()
    ip_reserved = IPAddress.query.filter_by(status='reserved').count()
    ip_disabled = IPAddress.query.filter_by(status='disabled').count()
    usage_percent = round(ip_allocated / ip_total * 100, 1) if ip_total else 0

    recent_subnets = Subnet.query.order_by(Subnet.created_at.desc()).limit(8).all()
    recent_ips = IPAddress.query.order_by(IPAddress.updated_at.desc()).limit(10).all()

    top_subnets = []
    for sub in Subnet.query.order_by(Subnet.id).all():
        usage = sub.ip_usage
        top_subnets.append((sub, usage))
    top_subnets.sort(key=lambda x: x[1]['allocated'], reverse=True)

    return render_template(
        'network/overview.html',
        subnet_count=subnet_count,
        vlan_count=vlan_count,
        ip_total=ip_total,
        ip_allocated=ip_allocated,
        ip_available=ip_available,
        ip_reserved=ip_reserved,
        ip_disabled=ip_disabled,
        usage_percent=usage_percent,
        recent_subnets=recent_subnets,
        recent_ips=recent_ips,
        top_subnets=top_subnets[:6],
    )


# ==================== 子网管理 ====================

@network_bp.route('/subnets')
@login_required
@permission_required('network:view')
def subnet_list():
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 20, type=int)
    q = request.args.get('q', '').strip()
    vlan_id = request.args.get('vlan_id', type=int)
    status = request.args.get('status', '').strip()

    query = Subnet.query
    if q:
        like = f'%{q}%'
        query = query.filter(db.or_(
            Subnet.name.like(like),
            Subnet.network.like(like),
            Subnet.gateway.like(like),
            Subnet.department.like(like),
            Subnet.tags.like(like),
        ))
    if vlan_id:
        query = query.filter(Subnet.vlan_id == vlan_id)
    if status:
        query = query.filter(Subnet.status == status)

    pagination = query.order_by(Subnet.network, Subnet.prefix_length).paginate(
        page=page, per_page=per_page, error_out=False)

    args = {'q': q, 'vlan_id': vlan_id or '', 'status': status, 'page': page}
    return render_template(
        'network/subnet_list.html',
        pagination=pagination,
        subnets=pagination.items,
        q=q,
        vlan_id=vlan_id,
        status=status,
        vlans=Vlan.query.order_by(Vlan.vlan_id).all(),
        args=args,
    )


@network_bp.route('/subnets/add', methods=['GET', 'POST'])
@login_required
@permission_required('network:edit')
def subnet_add():
    ctx = _subnet_generation_context()
    form = SubnetForm(data=ctx['preset']) if request.method == 'GET' else SubnetForm()
    if form.validate_on_submit():
        network = str(form.network.data).strip()
        prefix = form.prefix_length.data
        net = _parse_subnet(network, prefix)
        if not net:
            flash('网络地址或前缀长度无效，请输入合法的 IPv4 网段', 'danger')
            return render_template('network/subnet_form.html', form=form, title='新增子网', subnet=None, **ctx)
        overlap = _subnet_overlap(net)
        if overlap:
            flash(f'与已有子网 {overlap.name} ({overlap.cidr}) 重叠，请检查', 'danger')
            return render_template('network/subnet_form.html', form=form, title='新增子网', subnet=None, **ctx)

        subnet = Subnet(
            name=form.name.data.strip(),
            network=str(net.network_address),
            prefix_length=net.prefixlen,
            gateway=_normalize_ip(form.gateway.data),
            vlan_id=form.vlan_id.data if form.vlan_id.data and form.vlan_id.data > 0 else None,
            location_id=form.location_id.data if form.location_id.data and form.location_id.data > 0 else None,
            department=form.department.data.strip() if form.department.data else '',
            tags=form.tags.data.strip() if form.tags.data else '',
            description=form.description.data,
            status=form.status.data,
            ipam_enabled=form.ipam_enabled.data,
        )
        db.session.add(subnet)
        try:
            db.session.commit()
            _audit('create', 'subnet', subnet.id, f'新增子网: {subnet.name} ({subnet.cidr})')
            flash(f'子网 "{subnet.name} ({subnet.cidr})" 创建成功', 'success')
            if form.ipam_enabled.data and form.auto_generate.data:
                result = _generate_subnet_ips(subnet)
                msg = f'已生成 {result["created"]} 个 IP 记录'
                if result['limited']:
                    msg += f'（子网地址过多，仅生成前 {MAX_AUTO_GENERATE} 个）'
                flash(msg, 'info')
            session.pop('subnet_device_ids', None)
            return redirect(url_for('network.subnet_detail', id=subnet.id))
        except Exception as e:
            db.session.rollback()
            flash(f'创建失败: {e}', 'danger')
    return render_template('network/subnet_form.html', form=form, title='新增子网', subnet=None, **ctx)


@network_bp.route('/subnets/add/devices', methods=['POST'])
@login_required
@permission_required('network:edit')
def subnet_add_devices():
    """接收从设备列表生成子网的 POST 请求，将 device_ids 存入 session 后重定向到 GET 页面。
    避免数百个 device_ids 拼入 GET URL 导致超长 URI 被生产服务器拒绝。
    """
    device_ids = request.form.getlist('device_ids', type=int)
    session['subnet_device_ids'] = device_ids
    return redirect(url_for('network.subnet_add', candidate=0))


def _subnet_generation_context():
    """新增子网页的"从设备分组/设备生成"上下文（GET 预填数据）。
    device_ids 从 session 读取（POST 提交后存入），避免数百个 ID 拼入 URL。
    """
    groups = DeviceGroup.query.order_by(DeviceGroup.sort_order, DeviceGroup.name).all()
    all_devices = [
        d for d in Device.query.filter_by(is_decommissioned=False).order_by(Device.name).all()
        if _normalize_ip(d.management_ip) or _normalize_ip(d.ip_address)
    ]

    group_id = request.args.get('group_id', type=int)
    candidate_idx = request.args.get('candidate', 0, type=int)

    # device_ids 优先从 session 读取；group_id 模式下清除 session 缓存
    if group_id:
        session.pop('subnet_device_ids', None)
        device_ids = []
    else:
        device_ids = session.get('subnet_device_ids', [])

    source_desc = None
    group = None
    devices = []
    if group_id:
        group = DeviceGroup.query.get(group_id)
        if group:
            devices = _collect_group_devices(group)
            source_desc = f'设备分组「{group.name}」'
    elif device_ids:
        devices = Device.query.filter(Device.id.in_(device_ids)).all()
        source_desc = f'手动选择的 {len(devices)} 台设备'

    candidates = []
    preset = {}
    if devices:
        ip_list = _device_ip_pairs(devices)
        candidates = _cluster_subnet_candidates(ip_list)
        if candidates:
            preset = _build_generated_subnet_preset(
                candidates, candidate_idx, source_desc, group.name if group else '')

    return {
        'groups': groups,
        'all_devices': all_devices,
        'candidates': candidates,
        'source_desc': source_desc,
        'selected_group_id': group_id,
        'selected_device_ids': device_ids,
        'selected_candidate': candidate_idx,
        'preset': preset,
    }


@network_bp.route('/subnets/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@permission_required('network:edit')
def subnet_edit(id):
    subnet = Subnet.query.get_or_404(id)
    form = SubnetForm(obj=subnet)
    if form.validate_on_submit():
        network = str(form.network.data).strip()
        prefix = form.prefix_length.data
        net = _parse_subnet(network, prefix)
        if not net:
            flash('网络地址或前缀长度无效，请输入合法的 IPv4 网段', 'danger')
            return render_template('network/subnet_form.html', form=form, title='编辑子网', subnet=subnet)
        overlap = _subnet_overlap(net, exclude_id=subnet.id)
        if overlap:
            flash(f'与已有子网 {overlap.name} ({overlap.cidr}) 重叠，请检查', 'danger')
            return render_template('network/subnet_form.html', form=form, title='编辑子网', subnet=subnet)

        subnet.name = form.name.data.strip()
        subnet.network = str(net.network_address)
        subnet.prefix_length = net.prefixlen
        subnet.gateway = _normalize_ip(form.gateway.data)
        subnet.vlan_id = form.vlan_id.data if form.vlan_id.data and form.vlan_id.data > 0 else None
        subnet.location_id = form.location_id.data if form.location_id.data and form.location_id.data > 0 else None
        subnet.department = form.department.data.strip() if form.department.data else ''
        subnet.tags = form.tags.data.strip() if form.tags.data else ''
        subnet.description = form.description.data
        subnet.status = form.status.data
        subnet.ipam_enabled = form.ipam_enabled.data
        try:
            db.session.commit()
            _audit('update', 'subnet', subnet.id, f'编辑子网: {subnet.name} ({subnet.cidr})')
            flash(f'子网 "{subnet.name}" 更新成功', 'success')
            if form.ipam_enabled.data and form.auto_generate.data:
                result = _generate_subnet_ips(subnet)
                msg = f'已补生成 {result["created"]} 个 IP 记录（跳过 {result["skipped"]} 个已存在）'
                if result['limited']:
                    msg += f'（子网地址过多，仅生成前 {MAX_AUTO_GENERATE} 个）'
                flash(msg, 'info')
            return redirect(url_for('network.subnet_detail', id=subnet.id))
        except Exception as e:
            db.session.rollback()
            flash(f'更新失败: {e}', 'danger')
    return render_template('network/subnet_form.html', form=form, title='编辑子网', subnet=subnet)


@network_bp.route('/subnets/<int:id>/delete', methods=['POST'])
@login_required
@permission_required('network:edit')
def subnet_delete(id):
    subnet = Subnet.query.get_or_404(id)
    name, cidr = subnet.name, subnet.cidr
    ip_count = subnet.ip_addresses.count() if subnet.ip_addresses else 0
    try:
        db.session.delete(subnet)
        db.session.commit()
        _audit('delete', 'subnet', id, f'删除子网: {name} ({cidr})，连带删除 {ip_count} 条 IP 记录')
        flash(f'子网 "{name}" 已删除（含 {ip_count} 条 IP 记录）', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'删除失败: {e}', 'danger')
    return redirect(url_for('network.subnet_list'))


@network_bp.route('/subnets/<int:id>')
@login_required
@permission_required('network:view')
def subnet_detail(id):
    subnet = Subnet.query.get_or_404(id)
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 50, type=int)
    status_filter = request.args.get('status', '').strip()

    query = subnet.ip_addresses
    if status_filter:
        query = query.filter(IPAddress.status == status_filter)
    pagination = query.order_by(IPAddress.ip_address).paginate(page=page, per_page=per_page, error_out=False)

    return render_template(
        'network/subnet_detail.html',
        subnet=subnet,
        pagination=pagination,
        ips=pagination.items,
        status_filter=status_filter,
        args={'page': page, 'status': status_filter},
    )


@network_bp.route('/subnets/<int:id>/generate', methods=['POST'])
@login_required
@permission_required('network:edit')
def subnet_generate_ips(id):
    subnet = Subnet.query.get_or_404(id)
    result = _generate_subnet_ips(subnet)
    msg = f'新生成 {result["created"]} 个 IP（跳过 {result["skipped"]} 个已存在）'
    if result['limited']:
        msg += f'；子网共 {result["total"]} 个地址，超过上限仅生成前 {MAX_AUTO_GENERATE} 个'
    _audit('generate', 'subnet', id, f'批量生成 IP: {msg}')
    flash(msg, 'success' if result['created'] else 'info')
    return redirect(url_for('network.subnet_detail', id=subnet.id))


@network_bp.route('/subnets/<int:id>/sync-devices', methods=['POST'])
@login_required
@permission_required('network:edit')
def subnet_sync_devices(id):
    subnet = Subnet.query.get_or_404(id)
    result = _sync_devices_for_subnet(subnet)
    msg = f'匹配到 {result["matched"]} 台设备，新建 {result["created"]} 条、更新 {result["updated"]} 条 IP 记录'
    _audit('sync', 'subnet', id, f'从设备库同步 IP: {msg}')
    flash(msg, 'success')
    return redirect(url_for('network.subnet_detail', id=subnet.id))


@network_bp.route('/subnets/<int:id>/export')
@login_required
@permission_required('network:view')
def subnet_export_ips(id):
    subnet = Subnet.query.get_or_404(id)
    ips = subnet.ip_addresses.order_by(IPAddress.ip_address).all()
    output = _export_ips_to_excel(ips, sheet_title='子网IP')
    filename = f'{subnet.name}_{subnet.network}_{subnet.prefix_length}_IP地址.xlsx'
    return _send_excel(output, filename)


# ==================== IP 地址管理 ====================

@network_bp.route('/ips')
@login_required
@permission_required('network:view')
def ip_list():
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 50, type=int)
    subnet_id = request.args.get('subnet_id', type=int)
    vlan_id = request.args.get('vlan_id', type=int)
    status = request.args.get('status', '').strip()
    q = request.args.get('q', '').strip()

    query = IPAddress.query
    if subnet_id:
        query = query.filter(IPAddress.subnet_id == subnet_id)
    if vlan_id:
        query = query.join(Subnet, Subnet.id == IPAddress.subnet_id).filter(Subnet.vlan_id == vlan_id)
    if status:
        query = query.filter(IPAddress.status == status)
    if q:
        like = f'%{q}%'
        query = query.filter(db.or_(
            IPAddress.ip_address.like(like),
            IPAddress.hostname.like(like),
            IPAddress.owner.like(like),
            IPAddress.mac_address.like(like),
            IPAddress.department.like(like),
            IPAddress.description.like(like),
        ))

    pagination = query.order_by(IPAddress.ip_address).paginate(page=page, per_page=per_page, error_out=False)
    args = {'subnet_id': subnet_id or '', 'vlan_id': vlan_id or '', 'status': status, 'q': q, 'page': page}

    return render_template(
        'network/ip_list.html',
        pagination=pagination,
        ips=pagination.items,
        subnets=Subnet.query.order_by(Subnet.network).all(),
        vlans=Vlan.query.order_by(Vlan.vlan_id).all(),
        subnet_id=subnet_id,
        vlan_id=vlan_id,
        status=status,
        q=q,
        args=args,
        status_labels=IP_STATUS_LABELS,
    )


@network_bp.route('/ips/add', methods=['GET', 'POST'])
@login_required
@permission_required('network:edit')
def ip_add():
    form = IPAddressForm()
    preset_subnet = request.args.get('subnet_id', type=int)
    if preset_subnet and not form.is_submitted():
        form.subnet_id.data = preset_subnet
    if form.validate_on_submit():
        subnet = Subnet.query.get(form.subnet_id.data)
        if not subnet:
            flash('请选择有效的子网', 'danger')
            return render_template('network/ip_form.html', form=form, title='新增 IP 地址', ip=None)
        ip = _normalize_ip(form.ip_address.data)
        if not ip:
            flash('IP 地址格式无效', 'danger')
            return render_template('network/ip_form.html', form=form, title='新增 IP 地址', ip=None)
        try:
            if ipmod.ip_address(ip) not in _parse_subnet(subnet.network, subnet.prefix_length):
                flash(f'IP {ip} 不在子网 {subnet.cidr} 范围内', 'danger')
                return render_template('network/ip_form.html', form=form, title='新增 IP 地址', ip=None)
        except (TypeError, ValueError):
            flash('子网信息无效', 'danger')
            return render_template('network/ip_form.html', form=form, title='新增 IP 地址', ip=None)

        if IPAddress.query.filter_by(subnet_id=subnet.id, ip_address=ip).first():
            flash(f'IP {ip} 已存在于子网 {subnet.cidr} 中', 'danger')
            return render_template('network/ip_form.html', form=form, title='新增 IP 地址', ip=None)

        record = IPAddress(
            subnet_id=subnet.id,
            ip_address=ip,
            status=form.status.data,
            device_id=form.device_id.data if form.device_id.data and form.device_id.data > 0 else None,
            hostname=form.hostname.data.strip() if form.hostname.data else '',
            mac_address=form.mac_address.data.strip() if form.mac_address.data else '',
            owner=form.owner.data.strip() if form.owner.data else '',
            department=form.department.data.strip() if form.department.data else '',
            description=form.description.data,
        )
        db.session.add(record)
        try:
            db.session.commit()
            _audit('create', 'ip_address', record.id, f'新增 IP 记录: {ip} ({subnet.cidr})')
            flash(f'IP {ip} 添加成功', 'success')
            return redirect(url_for('network.ip_list', subnet_id=subnet.id))
        except Exception as e:
            db.session.rollback()
            flash(f'添加失败: {e}', 'danger')
    return render_template('network/ip_form.html', form=form, title='新增 IP 地址', ip=None)


@network_bp.route('/ips/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@permission_required('network:edit')
def ip_edit(id):
    record = IPAddress.query.get_or_404(id)
    form = IPAddressForm(obj=record)
    if form.validate_on_submit():
        subnet = Subnet.query.get(form.subnet_id.data)
        if not subnet:
            flash('请选择有效的子网', 'danger')
            return render_template('network/ip_form.html', form=form, title='编辑 IP 地址', ip=record)
        ip = _normalize_ip(form.ip_address.data)
        if not ip:
            flash('IP 地址格式无效', 'danger')
            return render_template('network/ip_form.html', form=form, title='编辑 IP 地址', ip=record)
        try:
            if ipmod.ip_address(ip) not in _parse_subnet(subnet.network, subnet.prefix_length):
                flash(f'IP {ip} 不在子网 {subnet.cidr} 范围内', 'danger')
                return render_template('network/ip_form.html', form=form, title='编辑 IP 地址', ip=record)
        except (TypeError, ValueError):
            flash('子网信息无效', 'danger')
            return render_template('network/ip_form.html', form=form, title='编辑 IP 地址', ip=record)

        dup = IPAddress.query.filter(
            IPAddress.subnet_id == subnet.id,
            IPAddress.ip_address == ip,
            IPAddress.id != record.id,
        ).first()
        if dup:
            flash(f'IP {ip} 已存在于子网 {subnet.cidr} 中', 'danger')
            return render_template('network/ip_form.html', form=form, title='编辑 IP 地址', ip=record)

        record.subnet_id = subnet.id
        record.ip_address = ip
        record.status = form.status.data
        record.device_id = form.device_id.data if form.device_id.data and form.device_id.data > 0 else None
        record.hostname = form.hostname.data.strip() if form.hostname.data else ''
        record.mac_address = form.mac_address.data.strip() if form.mac_address.data else ''
        record.owner = form.owner.data.strip() if form.owner.data else ''
        record.department = form.department.data.strip() if form.department.data else ''
        record.description = form.description.data
        try:
            db.session.commit()
            _audit('update', 'ip_address', record.id, f'编辑 IP 记录: {ip} ({subnet.cidr})')
            flash(f'IP {ip} 更新成功', 'success')
            return redirect(url_for('network.ip_list', subnet_id=subnet.id))
        except Exception as e:
            db.session.rollback()
            flash(f'更新失败: {e}', 'danger')
    return render_template('network/ip_form.html', form=form, title='编辑 IP 地址', ip=record)


@network_bp.route('/ips/<int:id>/delete', methods=['POST'])
@login_required
@permission_required('network:edit')
def ip_delete(id):
    record = IPAddress.query.get_or_404(id)
    ip = record.ip_address
    try:
        db.session.delete(record)
        db.session.commit()
        _audit('delete', 'ip_address', id, f'删除 IP 记录: {ip}')
        flash(f'IP {ip} 已删除', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'删除失败: {e}', 'danger')
    return redirect(request.referrer or url_for('network.ip_list'))


@network_bp.route('/ips/<int:id>/release', methods=['POST'])
@login_required
@permission_required('network:edit')
def ip_release(id):
    record = IPAddress.query.get_or_404(id)
    record.status = 'available'
    record.device_id = None
    record.hostname = ''
    record.mac_address = ''
    record.owner = ''
    try:
        db.session.commit()
        _audit('release', 'ip_address', id, f'释放 IP: {record.ip_address}')
        flash(f'IP {record.ip_address} 已释放', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'释放失败: {e}', 'danger')
    return redirect(request.referrer or url_for('network.ip_list'))


@network_bp.route('/ips/import', methods=['GET', 'POST'])
@login_required
@permission_required('network:edit')
def ip_import():
    subnets = Subnet.query.filter(Subnet.status == 'active').order_by(Subnet.network).all()
    preset_subnet = request.args.get('subnet_id', type=int)
    if request.method == 'POST':
        subnet_id = request.form.get('subnet_id', type=int)
        update_existing = request.form.get('update_existing') == 'on'
        auto_link = request.form.get('auto_link') == 'on'
        subnet = Subnet.query.get(subnet_id) if subnet_id else None
        if not subnet:
            flash('请选择目标子网', 'danger')
            return render_template('network/ip_import.html', subnets=subnets, preset_subnet=preset_subnet)

        net = _parse_subnet(subnet.network, subnet.prefix_length)
        rows = []
        file = request.files.get('file')
        text = (request.form.get('ip_text') or '').strip()
        if file and file.filename:
            try:
                rows = _read_import_file(file)
            except ValueError as e:
                flash(str(e), 'danger')
                return render_template('network/ip_import.html', subnets=subnets, preset_subnet=preset_subnet)
        elif text:
            for line in text.splitlines():
                line = line.strip()
                if line:
                    rows.append({'ip': line.split()[0] if line.split() else line})
        else:
            flash('请上传文件或粘贴 IP 列表', 'danger')
            return render_template('network/ip_import.html', subnets=subnets, preset_subnet=preset_subnet)

        created = updated = skipped = errors = 0
        device_cache = {d.name: d for d in Device.query.filter_by(is_decommissioned=False).all()}
        ip_device_map = {}
        if auto_link:
            for d in Device.query.filter_by(is_decommissioned=False).all():
                for dev_ip in (d.management_ip, d.ip_address):
                    if dev_ip:
                        ip_device_map.setdefault(dev_ip, d)
        ip_cache = {ip.ip_address: ip for ip in IPAddress.query.filter_by(subnet_id=subnet.id).all()}

        for row in rows:
            ip = _normalize_ip(row.get('ip'))
            if not ip:
                errors += 1
                continue
            if net is None or ipmod.ip_address(ip) not in net:
                errors += 1
                continue

            device = None
            device_name = (row.get('device_name') or '').strip()
            if device_name:
                device = device_cache.get(device_name)
            if not device and auto_link:
                device = ip_device_map.get(ip)

            status = (row.get('status') or '').strip()
            if not status:
                status = 'allocated' if device else 'available'
            status = status if status in IP_STATUS_LABELS else 'available'

            existing = ip_cache.get(ip)
            if existing:
                if update_existing:
                    if device:
                        existing.status = 'allocated'
                        existing.device_id = device.id
                    else:
                        existing.status = status
                    existing.hostname = row.get('hostname') or existing.hostname or ''
                    existing.mac_address = row.get('mac') or existing.mac_address or ''
                    existing.owner = row.get('owner') or existing.owner or ''
                    existing.department = row.get('department') or existing.department or ''
                    if row.get('description'):
                        existing.description = row.get('description')
                    updated += 1
                else:
                    skipped += 1
                continue

            record = IPAddress(
                subnet_id=subnet.id,
                ip_address=ip,
                status=status,
                device_id=device.id if device else None,
                hostname=(row.get('hostname') or '').strip(),
                mac_address=(row.get('mac') or '').strip(),
                owner=(row.get('owner') or '').strip(),
                department=(row.get('department') or '').strip(),
                description=(row.get('description') or '').strip() or ('自动关联设备' if device else ''),
            )
            db.session.add(record)
            ip_cache[ip] = record
            created += 1

        try:
            db.session.commit()
        except Exception as e:
            db.session.rollback()
            flash(f'导入失败: {e}', 'danger')
            return render_template('network/ip_import.html', subnets=subnets, preset_subnet=preset_subnet)

        _audit('import', 'ip_address', subnet.id,
               f'导入 IP 到 {subnet.cidr}: 新建 {created}，更新 {updated}，跳过 {skipped}，错误 {errors}')
        flash(f'导入完成：新建 {created} 条，更新 {updated} 条，跳过 {skipped} 条，无效 {errors} 条', 'success')
        return redirect(url_for('network.subnet_detail', id=subnet.id))

    return render_template('network/ip_import.html', subnets=subnets, preset_subnet=preset_subnet)


def _read_import_file(file):
    """读取 Excel/CSV 导入文件，返回规范化行字典列表。"""
    import pandas as pd
    filename = (file.filename or '').lower()
    try:
        if filename.endswith('.csv'):
            df = pd.read_csv(file, dtype=str, encoding='utf-8-sig')
        else:
            df = pd.read_excel(file, dtype=str)
    except Exception as e:
        raise ValueError(f'文件解析失败: {e}')

    df = df.fillna('')
    col_map = {}
    for col in df.columns:
        key = str(col).strip().lower()
        if key in ('ip', 'ip地址', 'ip_address', 'ip地址/主机名'):
            col_map['ip'] = col
        elif key in ('状态', 'status'):
            col_map['status'] = col
        elif key in ('主机名', 'hostname', '设备名称', 'device_name'):
            col_map['device_name'] = col
            col_map.setdefault('hostname', col)
        elif key in ('mac', 'mac地址', 'mac_address'):
            col_map['mac'] = col
        elif key in ('使用人', 'owner'):
            col_map['owner'] = col
        elif key in ('部门', 'department'):
            col_map['department'] = col
        elif key in ('用途', '备注', '描述', 'description'):
            col_map['description'] = col
    if 'ip' not in col_map:
        raise ValueError('未找到 IP 地址列（支持列名: IP地址 / ip / ip_address）')

    rows = []
    for _, row in df.iterrows():
        item = {}
        for field, col in col_map.items():
            value = row.get(col)
            item[field] = str(value).strip() if value is not None else ''
        rows.append(item)
    return rows


@network_bp.route('/ips/export')
@login_required
@permission_required('network:view')
def ip_export():
    subnet_id = request.args.get('subnet_id', type=int)
    vlan_id = request.args.get('vlan_id', type=int)
    status = request.args.get('status', '').strip()
    q = request.args.get('q', '').strip()

    query = IPAddress.query
    if subnet_id:
        query = query.filter(IPAddress.subnet_id == subnet_id)
    if vlan_id:
        query = query.join(Subnet, Subnet.id == IPAddress.subnet_id).filter(Subnet.vlan_id == vlan_id)
    if status:
        query = query.filter(IPAddress.status == status)
    if q:
        like = f'%{q}%'
        query = query.filter(db.or_(
            IPAddress.ip_address.like(like),
            IPAddress.hostname.like(like),
            IPAddress.owner.like(like),
        ))
    ips = query.order_by(IPAddress.ip_address).all()
    output = _export_ips_to_excel(ips, sheet_title='IP地址')
    return _send_excel(output, 'IP地址导出.xlsx')


# ==================== VLAN 管理 ====================

@network_bp.route('/vlans')
@login_required
@permission_required('network:view')
def vlan_list():
    q = request.args.get('q', '').strip()
    query = Vlan.query
    if q:
        like = f'%{q}%'
        query = query.filter(db.or_(
            Vlan.name.like(like),
            Vlan.description.like(like),
            Vlan.vlan_type.like(like),
        ))
    vlans = query.order_by(Vlan.vlan_id).all()
    return render_template('network/vlan_list.html', vlans=vlans, q=q)


@network_bp.route('/vlans/add', methods=['GET', 'POST'])
@login_required
@permission_required('network:edit')
def vlan_add():
    form = VlanForm()
    if form.validate_on_submit():
        if Vlan.query.filter_by(vlan_id=form.vlan_id.data).first():
            flash(f'VLAN {form.vlan_id.data} 已存在', 'danger')
            return render_template('network/vlan_form.html', form=form, title='新增 VLAN', vlan=None)
        vlan = Vlan(
            vlan_id=form.vlan_id.data,
            name=form.name.data.strip(),
            vlan_type=form.vlan_type.data,
            gateway=_normalize_ip(form.gateway.data),
            description=form.description.data,
            status=form.status.data,
        )
        db.session.add(vlan)
        try:
            db.session.commit()
            _audit('create', 'vlan', vlan.id, f'新增 VLAN: {vlan.vlan_id} {vlan.name}')
            flash(f'VLAN {vlan.vlan_id} ({vlan.name}) 创建成功', 'success')
            return redirect(url_for('network.vlan_list'))
        except Exception as e:
            db.session.rollback()
            flash(f'创建失败: {e}', 'danger')
    return render_template('network/vlan_form.html', form=form, title='新增 VLAN', vlan=None)


@network_bp.route('/vlans/<int:id>/edit', methods=['GET', 'POST'])
@login_required
@permission_required('network:edit')
def vlan_edit(id):
    vlan = Vlan.query.get_or_404(id)
    form = VlanForm(obj=vlan)
    if form.validate_on_submit():
        dup = Vlan.query.filter(Vlan.vlan_id == form.vlan_id.data, Vlan.id != vlan.id).first()
        if dup:
            flash(f'VLAN {form.vlan_id.data} 已被其他记录使用', 'danger')
            return render_template('network/vlan_form.html', form=form, title='编辑 VLAN', vlan=vlan)
        vlan.vlan_id = form.vlan_id.data
        vlan.name = form.name.data.strip()
        vlan.vlan_type = form.vlan_type.data
        vlan.gateway = _normalize_ip(form.gateway.data)
        vlan.description = form.description.data
        vlan.status = form.status.data
        try:
            db.session.commit()
            _audit('update', 'vlan', vlan.id, f'编辑 VLAN: {vlan.vlan_id} {vlan.name}')
            flash(f'VLAN {vlan.vlan_id} 更新成功', 'success')
            return redirect(url_for('network.vlan_list'))
        except Exception as e:
            db.session.rollback()
            flash(f'更新失败: {e}', 'danger')
    return render_template('network/vlan_form.html', form=form, title='编辑 VLAN', vlan=vlan)


@network_bp.route('/vlans/<int:id>/delete', methods=['POST'])
@login_required
@permission_required('network:edit')
def vlan_delete(id):
    vlan = Vlan.query.get_or_404(id)
    name, vid = vlan.name, vlan.vlan_id
    if vlan.subnets.count():
        flash(f'VLAN {vid} 下仍有 {vlan.subnets.count()} 个子网，请先解除关联', 'danger')
        return redirect(url_for('network.vlan_list'))
    try:
        db.session.delete(vlan)
        db.session.commit()
        _audit('delete', 'vlan', id, f'删除 VLAN: {vid} {name}')
        flash(f'VLAN {vid} ({name}) 已删除', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'删除失败: {e}', 'danger')
    return redirect(url_for('network.vlan_list'))


# ==================== AJAX ====================

@network_bp.route('/api/subnets/<int:id>/ips')
@login_required
@permission_required('network:view')
def api_subnet_ips(id):
    subnet = Subnet.query.get_or_404(id)
    status = request.args.get('status', '').strip()
    query = subnet.ip_addresses
    if status:
        query = query.filter(IPAddress.status == status)
    ips = query.order_by(IPAddress.ip_address).limit(2000).all()
    return jsonify({
        'subnet': subnet.to_dict(),
        'ips': [ip.to_dict() for ip in ips],
    })


@network_bp.route('/api/vlans/<int:id>/subnets')
@login_required
@permission_required('network:view')
def api_vlan_subnets(id):
    vlan = Vlan.query.get_or_404(id)
    return jsonify({
        'vlan': vlan.to_dict(),
        'subnets': [s.to_dict() for s in vlan.subnets.all()],
    })

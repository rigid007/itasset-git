# excel_utils.py
"""Excel import/export utilities for locations, devices, and templates."""
import io
import openpyxl
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO

import logging
logger = logging.getLogger(__name__)


def generate_location_template():
    """生成位置模板Excel文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "位置模板"
    headers = ['名称', '描述', '地址', '联系人', '联系电话', '备注']
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    example = ['核心机房', '主数据中心', '北京市海淀区中关村大街1号B1层', '张三', '13800138001', '24小时监控']
    for col, val in enumerate(example, 1):
        ws.cell(row=2, column=col, value=val)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_locations_to_excel(locations):
    """导出位置数据到Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "位置列表"
    headers = ["位置名称", "描述", "地址", "联系人", "联系电话", "备注"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    for row, loc in enumerate(locations, 2):
        ws.cell(row=row, column=1, value=loc.name or "")
        ws.cell(row=row, column=2, value=loc.description or "")
        ws.cell(row=row, column=3, value=loc.address or "")
        ws.cell(row=row, column=4, value=loc.contact_person or "")
        ws.cell(row=row, column=5, value=loc.contact_phone or "")
        ws.cell(row=row, column=6, value=loc.notes or "")
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def import_locations_from_excel(file_path, db, Location, log_operation):
    """从Excel导入位置数据"""
    df = pd.read_excel(file_path)
    required = ['名称']
    for col in required:
        if col not in df.columns:
            raise ValueError(f"缺少列: {col}")
    success = 0
    errors = []
    for idx, row in df.iterrows():
        name = str(row['名称']).strip()
        if not name:
            errors.append(f"第{idx+2}行: 名称为空")
            continue
        if Location.query.filter_by(name=name).first():
            errors.append(f"第{idx+2}行: 位置已存在")
            continue
        loc = Location(
            name=name,
            description=str(row.get('描述', '')).strip() if pd.notna(row.get('描述')) else None,
            address=str(row.get('地址', '')).strip() if pd.notna(row.get('地址')) else None,
            contact_person=str(row.get('联系人', '')).strip() if pd.notna(row.get('联系人')) else None,
            contact_phone=str(row.get('联系电话', '')).strip() if pd.notna(row.get('联系电话')) else None,
            notes=str(row.get('备注', '')).strip() if pd.notna(row.get('备注')) else None
        )
        db.session.add(loc)
        success += 1
    db.session.commit()
    return {'success': True, 'success_count': success, 'error_count': len(errors), 'error_messages': errors[:10]}


def export_devices_to_excel(devices):
    """导出设备数据到Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = '设备列表'
    headers = ['设备名称', '设备类型', '型号', '序列号', '资产编号', '机柜名称', 'U位',
               'U高度', 'IP地址', 'MAC地址', '状态', '安装时间', '位置名称', '联系人', '联系电话']
    ws.append(headers)
    for dev in devices:
        ws.append([
            dev.name, dev.device_type, dev.model, dev.serial_number, dev.asset_number,
            dev.cabinet.name if dev.cabinet else '', dev.u_position, dev.u_height,
            dev.ip_address, dev.mac_address, dev.status,
            dev.installed_at.strftime('%Y-%m-%d') if dev.installed_at else '',
            dev.location.name if dev.location else '', dev.contact_person, dev.contact_phone
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_device_template():
    """生成设备导入模板"""
    wb = Workbook()
    ws = wb.active
    ws.title = '设备导入模板'
    headers = ['设备名称（必填）', '设备类型', '型号', '序列号', '资产编号',
               '机柜名称（必填）', 'U位', 'U高度', 'IP地址', '状态']
    ws.append(headers)
    ws.append(['服务器01', 'server', 'DELL R750', 'SN123456', 'ASSET789',
               '机柜A01', 1, 2, '192.168.1.10', 'online'])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def import_devices_from_excel(filepath, db, Device, Cabinet, Location, log_operation):
    """从Excel导入设备数据"""
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    success = 0
    errors = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[0] if row[0] else None
        device_type = row[1] if row[1] else None
        model = row[2] if row[2] else None
        serial_number = row[3] if row[3] else None
        asset_number = row[4] if row[4] else None
        cabinet_name = row[5] if row[5] else None
        u_position = row[6] if row[6] else None
        u_height = row[7] if row[7] else 1
        ip_address = row[8] if row[8] else None
        status = row[9] if row[9] else 'online'
        if not name or not cabinet_name:
            errors.append(f"第{row[0].row if hasattr(row, 'row') else '?'}行: 名称或机柜为空")
            continue
        cabinet = Cabinet.query.filter_by(name=cabinet_name).first()
        if not cabinet:
            errors.append(f"机柜不存在: {cabinet_name}")
            continue
        if u_position and Device.query.filter_by(cabinet_id=cabinet.id, u_position=u_position).first():
            errors.append(f"U位冲突: {u_position}")
            continue
        dev = Device(
            name=name, device_type=device_type, model=model, serial_number=serial_number,
            asset_number=asset_number, cabinet_id=cabinet.id, location_id=cabinet.location_id,
            u_position=u_position, u_height=u_height, ip_address=ip_address, status=status,
            installed_at=datetime.utcnow(), created_at=datetime.utcnow(), updated_at=datetime.utcnow()
        )
        db.session.add(dev)
        success += 1
    db.session.commit()
    return {'success': True, 'success_count': success, 'error_count': len(errors), 'errors': errors[:10]}
